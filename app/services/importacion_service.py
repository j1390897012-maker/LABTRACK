"""Importación masiva de equipos y estudiantes desde un archivo .xlsx o .csv.

Reglas:
- Todo-o-nada: si alguna fila tiene error no se guarda nada y se devuelve
  la lista de errores con su número de fila (como se ve en Excel).
- Se puede repetir la carga: lo que ya existe (mismo `codigo` / misma
  `matricula`) se actualiza en vez de duplicarse.
- `dry_run=True` valida y cuenta sin guardar.
"""

import csv
import io
import re
import unicodedata
from collections.abc import Iterable
from datetime import datetime
from typing import Any

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.repositories.equipo_repository import EquipoRepository
from app.repositories.estudiante_repository import EstudianteRepository
from app.repositories.falla_repository import FallaRepository
from app.schemas.importacion import ErrorFila, ImportacionResponse

MAX_BYTES = 2 * 1024 * 1024
MAX_FILAS = 5000
TIPO_POR_DEFECTO = "Sin clasificar"

# Estados que se pueden cargar por archivo. "Prestado" no: un préstamo
# necesita estudiante y sesión, se genera escaneando.
ESTADOS_VALIDOS = {
    "disponible": "Disponible",
    "en revision": "En revisión",
    "baja": "Baja",
    # Vocabulario de la hoja de inventario del laboratorio (resultado de la
    # revisión física). Solo lo que se sabe que funciona queda disponible.
    "funciona bien": "Disponible",
    "funciona parcialmente": "En revisión",
    "no funciona": "En revisión",
    "sin revisar": "En revisión",
}

# Encabezados aceptados (ya normalizados) -> campo interno
ALIAS_EQUIPOS = {
    "codigo": "codigo",
    "codigo_lab": "codigo",
    "inventario_lab": "codigo",
    "inventario_laboratorio": "codigo",
    "no_inventario_lab": "codigo",
    "numero_inventario_lab": "codigo",
    "numero_de_inventario_del_lab": "codigo",
    "numero_inventario_laboratorio": "codigo",
    "numero_de_inventario_del_laboratorio": "codigo",
    "numero_inventario": "numero_inventario",
    "n_inventario": "numero_inventario",
    "no_inventario": "numero_inventario",
    "inventario": "numero_inventario",
    "numero_de_inventario": "numero_inventario",
    "marca": "marca",
    "modelo": "modelo",
    "tipo": "tipo",
    "equipo": "tipo",
    "estado": "estado",
    "observaciones": "observaciones",
    "observacion": "observaciones",
    "fecha": "fecha",
}
ALIAS_ESTUDIANTES = {
    "matricula": "matricula",
    "nombre": "nombre",
    "nombre_completo": "nombre",
    "estudiante": "nombre",
}


def _sin_acentos(texto: str) -> str:
    n = unicodedata.normalize("NFD", texto)
    return "".join(c for c in n if unicodedata.category(c) != "Mn")


def _encabezado(valor: Any) -> str:
    t = _sin_acentos(str(valor or "")).lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def _celda(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    return str(valor).strip()


def _parsear_fecha(texto: str) -> datetime | None:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, fmt)
        except ValueError:
            continue
    return None


def _filas_xlsx(contenido: bytes) -> list[list[Any]]:
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = wb.worksheets[0]
        return [list(r) for r in hoja.iter_rows(values_only=True)]
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se pudo leer el archivo Excel (.xlsx).",
        ) from exc


def _filas_csv(contenido: bytes) -> list[list[Any]]:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")  # CSV guardado por Excel en español
    muestra = texto[:2048]
    delim = ";" if muestra.count(";") > muestra.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(texto), delimiter=delim)]


def leer_tabla(
    nombre_archivo: str, contenido: bytes, alias: dict[str, str]
) -> tuple[dict[int, str], list[tuple[int, dict[str, str]]]]:
    """Devuelve (columnas, filas) donde filas = [(nº de fila Excel, datos)].
    Ignora filas completamente vacías."""
    if len(contenido) > MAX_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="El archivo excede 2 MB.",
        )
    nombre = (nombre_archivo or "").lower()
    if nombre.endswith(".xlsx"):
        crudo = _filas_xlsx(contenido)
    elif nombre.endswith(".csv"):
        crudo = _filas_csv(contenido)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Formato no soportado: sube un archivo .xlsx o .csv.",
        )
    if not crudo:
        raise HTTPException(status_code=400, detail="El archivo está vacío.")

    columnas: dict[int, str] = {}
    for i, h in enumerate(crudo[0]):
        campo = alias.get(_encabezado(h))
        if campo and campo not in columnas.values():
            columnas[i] = campo

    filas: list[tuple[int, dict[str, str]]] = []
    for n, fila in enumerate(crudo[1:], start=2):
        datos = {
            campo: _celda(fila[i]) if i < len(fila) else ""
            for i, campo in columnas.items()
        }
        if any(datos.values()):
            filas.append((n, datos))
    if len(filas) > MAX_FILAS:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"El archivo tiene más de {MAX_FILAS} filas.",
        )
    return columnas, filas


def _exigir_columnas(columnas: dict[int, str], requeridas: Iterable[str]) -> None:
    faltan = [c for c in requeridas if c not in columnas.values()]
    if faltan:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Faltan columnas obligatorias: " + ", ".join(faltan) + ".",
        )


class ImportacionService:
    def __init__(self) -> None:
        self.repo_equipo = EquipoRepository()
        self.repo_estudiante = EstudianteRepository()
        self.repo_falla = FallaRepository()

    def importar_equipos(
        self, db: Session, nombre_archivo: str, contenido: bytes, dry_run: bool
    ) -> ImportacionResponse:
        columnas, filas = leer_tabla(nombre_archivo, contenido, ALIAS_EQUIPOS)
        _exigir_columnas(columnas, ["codigo"])

        errores: list[ErrorFila] = []
        creados = actualizados = fallas = 0
        vistos_codigo: set[str] = set()
        vistos_inv: dict[str, str] = {}
        tipos: dict[str, int] = {}

        def err(n: int, valor: str | None, motivo: str) -> None:
            errores.append(ErrorFila(fila=n, valor=valor or None, motivo=motivo))

        for n, d in filas:
            codigo = d.get("codigo", "")
            if not codigo and not any(
                d.get(k) for k in ("numero_inventario", "marca", "modelo", "tipo")
            ):
                continue  # fila de plantilla vacía (a veces solo trae estado)
            if not codigo:
                err(n, None, "Falta el número de inventario del laboratorio.")
                continue
            if len(codigo) > 50:
                err(n, codigo, "El código excede 50 caracteres.")
                continue
            if codigo in vistos_codigo:
                err(n, codigo, "Código repetido dentro del archivo.")
                continue
            vistos_codigo.add(codigo)

            inv = d.get("numero_inventario") or None
            if inv:
                if len(inv) > 50:
                    err(n, codigo, "El número de inventario excede 50 caracteres.")
                    continue
                if inv in vistos_inv:
                    err(n, codigo, f"Inventario '{inv}' repetido en {vistos_inv[inv]}.")
                    continue
                vistos_inv[inv] = codigo
            marca = d.get("marca") or None
            modelo = d.get("modelo") or None
            if (marca and len(marca) > 100) or (modelo and len(modelo) > 100):
                err(n, codigo, "Marca o modelo exceden 100 caracteres.")
                continue

            estado_txt = d.get("estado", "")
            estado = None
            if estado_txt:
                clave = _sin_acentos(estado_txt).lower().strip()
                estado = ESTADOS_VALIDOS.get(clave)
                if estado is None:
                    err(
                        n,
                        codigo,
                        f"Estado '{estado_txt}' no válido (usa Disponible, "
                        "En revisión o Baja).",
                    )
                    continue

            obs = d.get("observaciones") or None
            fecha = None
            if d.get("fecha"):
                fecha = _parsear_fecha(d["fecha"])
                if fecha is None:
                    err(n, codigo, f"Fecha '{d['fecha']}' no válida (AAAA-MM-DD).")
                    continue

            if inv:
                dueno = self.repo_equipo.get_by_numero_inventario(db, inv)
                if dueno is not None and dueno.codigo != codigo:
                    err(
                        n,
                        codigo,
                        f"El inventario '{inv}' ya pertenece a {dueno.codigo}.",
                    )
                    continue

            tipo_nombre = d.get("tipo") or None
            tipo_id = None
            if tipo_nombre:
                if tipo_nombre not in tipos:
                    tipo = self.repo_equipo.get_tipo_by_nombre(db, tipo_nombre)
                    if tipo is None:
                        tipo = self.repo_equipo.crear_tipo_sin_commit(db, tipo_nombre)
                    tipos[tipo_nombre] = tipo.id
                tipo_id = tipos[tipo_nombre]

            existente = self.repo_equipo.get_by_codigo(db, codigo)
            if existente is None:
                if tipo_id is None:
                    if TIPO_POR_DEFECTO not in tipos:
                        t = self.repo_equipo.get_tipo_by_nombre(db, TIPO_POR_DEFECTO)
                        if t is None:
                            t = self.repo_equipo.crear_tipo_sin_commit(
                                db, TIPO_POR_DEFECTO
                            )
                        tipos[TIPO_POR_DEFECTO] = t.id
                    tipo_id = tipos[TIPO_POR_DEFECTO]
                existente = self.repo_equipo.crear_sin_commit(
                    db, codigo, tipo_id, estado or "Disponible", inv, marca, modelo
                )
                creados += 1
            else:
                if estado and estado != existente.estado:
                    if existente.estado == "Prestado":
                        err(
                            n,
                            codigo,
                            "El equipo está prestado: registra la "
                            "devolución primero.",
                        )
                        continue
                    existente.estado = estado
                if tipo_id is not None:
                    existente.tipo_equipo_id = tipo_id
                if inv:
                    existente.numero_inventario = inv
                if marca:
                    existente.marca = marca
                if modelo:
                    existente.modelo = modelo
                db.flush()
                actualizados += 1

            # Las observaciones del inventario son el reporte de fallas inicial.
            if obs and not self.repo_falla.existe_descripcion(db, existente.id, obs):
                self.repo_falla.crear_sin_commit(db, existente.id, obs, fecha)
                fallas += 1

        total = creados + actualizados + len(errores)
        return self._cerrar(
            db, total, creados, actualizados, errores, dry_run, fallas
        )

    def importar_estudiantes(
        self, db: Session, nombre_archivo: str, contenido: bytes, dry_run: bool
    ) -> ImportacionResponse:
        columnas, filas = leer_tabla(nombre_archivo, contenido, ALIAS_ESTUDIANTES)
        _exigir_columnas(columnas, ["matricula", "nombre"])

        errores: list[ErrorFila] = []
        creados = actualizados = 0
        vistas: set[str] = set()

        for n, d in filas:
            matricula, nombre = d.get("matricula", ""), d.get("nombre", "")
            if not matricula or not nombre:
                errores.append(
                    ErrorFila(
                        fila=n,
                        valor=matricula or None,
                        motivo="Faltan matrícula o nombre.",
                    )
                )
                continue
            if len(matricula) > 50 or len(nombre) > 150:
                errores.append(
                    ErrorFila(
                        fila=n,
                        valor=matricula,
                        motivo="Matrícula o nombre demasiado largos.",
                    )
                )
                continue
            if matricula in vistas:
                errores.append(
                    ErrorFila(
                        fila=n,
                        valor=matricula,
                        motivo="Matrícula repetida dentro del archivo.",
                    )
                )
                continue
            vistas.add(matricula)

            existente = self.repo_estudiante.get_by_matricula(db, matricula)
            if existente is None:
                self.repo_estudiante.crear_sin_commit(db, nombre, matricula)
                creados += 1
            else:
                existente.nombre = nombre  # no se toca uid_rfid ni historial
                db.flush()
                actualizados += 1

        return self._cerrar(db, len(filas), creados, actualizados, errores, dry_run)

    @staticmethod
    def _cerrar(
        db: Session,
        total: int,
        creados: int,
        actualizados: int,
        errores: list[ErrorFila],
        dry_run: bool,
        fallas: int = 0,
    ) -> ImportacionResponse:
        aplicado = not errores and not dry_run
        if aplicado:
            db.commit()
        else:
            db.rollback()
        return ImportacionResponse(
            total_filas=total,
            creados=creados,
            actualizados=actualizados,
            fallas_creadas=fallas,
            errores=errores,
            aplicado=aplicado,
            dry_run=dry_run,
        )
